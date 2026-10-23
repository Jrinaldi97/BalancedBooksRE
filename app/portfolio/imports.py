"""Deterministic CSV/XLSX/XLS normalization; no guessed dates or account identities."""

import csv
import hashlib
import io
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from app.portfolio.service import fail, Problem


def parse_bank(content, filename, mapping, date_format="%Y-%m-%d", sheet=None):
    if len(content) > 25 * 1024 * 1024:
        fail("file_size", "Bank import limit is 25 MB per file")
    if (
        not isinstance(mapping, dict)
        or not mapping.get("date")
        or not (mapping.get("amount") or mapping.get("credit") or mapping.get("debit"))
    ):
        fail(
            "invalid_mapping",
            "Map a date and an amount or debit/credit column explicitly",
        )
    if any(
        key
        not in {
            "date",
            "amount",
            "debit",
            "credit",
            "description",
            "reference",
            "external_id",
        }
        or not isinstance(column, str)
        or not column
        for key, column in mapping.items()
    ):
        fail("invalid_mapping", "Use supported fields and explicit column names")
    source_sheet = None
    physical_rows = []
    if filename.lower().endswith(".xls"):
        from app.portfolio.legacy_xls import xls_rows

        source_sheet, raw_rows = xls_rows(content, sheet)
        headers = raw_rows[0] if raw_rows else []
        rows = [dict(zip([str(h or "") for h in headers], row)) for row in raw_rows[1:]]
        physical_rows = list(range(2, len(rows) + 2))
    elif filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        if not content.startswith(b"PK\x03\x04"):
            fail("file_type", "The .xlsx filename must contain an XLSX ZIP workbook")
        from zipfile import BadZipFile
        from openpyxl.utils.exceptions import InvalidFileException

        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        except (
            BadZipFile,
            InvalidFileException,
            KeyError,
            ValueError,
            OSError,
            SyntaxError,
        ):
            fail("invalid_workbook", "The XLSX workbook is corrupt or unsupported")
        try:
            if sheet is None and len(book.sheetnames) != 1:
                fail(
                    "sheet_required",
                    "Select an explicit sheet from a multi-sheet export",
                )
            ws = book[sheet] if sheet else book.active
            iterator = ws.iter_rows(values_only=True)
            headers = next(iterator, None)
            if not headers:
                fail("empty_file", "Spreadsheet has no header row")
            rows = []
            for row in iterator:
                if len(rows) >= 20000:
                    fail(
                        "row_limit",
                        "Split exports into files of at most 20,000 data rows",
                    )
                if any(
                    isinstance(value, str) and value.startswith("=") for value in row
                ):
                    fail(
                        "formula_cell", "XLSX contains formulas; export values instead"
                    )
                rows.append(dict(zip([str(h or "") for h in headers], row)))
            physical_rows = list(range(2, len(rows) + 2))
            source_sheet = ws.title
        except Problem:
            raise
        except (
            BadZipFile,
            InvalidFileException,
            KeyError,
            ValueError,
            OSError,
            SyntaxError,
        ):
            fail(
                "invalid_workbook",
                "The XLSX workbook is corrupt or has no selected sheet",
            )
        finally:
            book.close()
    elif filename.lower().endswith(".csv"):
        if (
            content.startswith((b"PK\x03\x04", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"))
            or b"\0" in content
        ):
            fail("file_type", "The .csv filename must contain a text CSV export")
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1252")
        reader = csv.reader(io.StringIO(text))
        headers = next(reader, [])
        rows = []
        while True:
            physical_start = reader.line_num + 1
            row = next(reader, None)
            if row is None:
                break
            if not row:
                continue
            if len(row) > len(headers):
                fail(
                    "invalid_row",
                    f"CSV row {physical_start} has more values than headers",
                )
            if len(rows) >= 20000:
                fail(
                    "row_limit", "Split exports into files of at most 20,000 data rows"
                )
            rows.append(dict(zip(headers, row)))
            physical_rows.append(physical_start)
    else:
        fail(
            "file_type",
            "Import CSV, XLSX or BIFF XLS exports; link PDF statements as supporting documents",
        )
    header_names = [str(h or "") for h in headers]
    named_headers = [h for h in header_names if h]
    if not headers or len(named_headers) != len(set(named_headers)):
        fail("invalid_header", "Use one header row with unique column names")
    if any(column not in header_names for column in mapping.values()):
        fail("missing_column", "A mapped column is missing from the header row")
    if len(rows) > 20000:
        fail("row_limit", "Split exports into files of at most 20,000 rows")

    def money(value):
        text = str(value or "0").strip().replace("$", "").replace(",", "")
        if text.startswith("(") and text.endswith(")"):
            text = "-" + text[1:-1]
        try:
            number = Decimal(text)
        except InvalidOperation:
            fail("invalid_amount", f"Invalid bank amount: {text[:50]}")
        if not number.is_finite() or number * 100 != (number * 100).to_integral_value():
            fail("invalid_amount", "Bank amounts must be finite and exact to cents")
        return int(number * 100)

    result = []
    for number, row in zip(physical_rows, rows, strict=True):
        if all(v is None or str(v).strip() == "" for v in row.values()):
            continue

        def cell(key, default=""):
            column = mapping.get(key)
            if column and column not in row:
                fail("missing_column", f"Column {column} is missing")
            value = row.get(column, default)
            if isinstance(value, str) and value.startswith("="):
                fail(
                    "formula_cell",
                    f"Row {number} contains an unevaluated formula; export values instead",
                )
            return value

        raw = cell("date")
        try:
            day = (
                raw.date()
                if isinstance(raw, datetime)
                else (
                    raw
                    if isinstance(raw, date)
                    else datetime.strptime(str(raw).strip(), date_format).date()
                )
            )
        except ValueError:
            fail("invalid_date", f"Row {number}: date does not match {date_format}")
        amount = (
            money(cell("amount"))
            if mapping.get("amount")
            else money(cell("credit")) - money(cell("debit"))
        )
        result.append(
            {
                "physical_row": number,
                "source_sheet": source_sheet,
                "date": day.isoformat(),
                "amount_cents": amount,
                "description": str(cell("description") or ""),
                "reference": str(cell("reference") or ""),
                "external_id": (
                    str(cell("external_id"))
                    if cell("external_id") not in ("", None)
                    else None
                ),
            }
        )
    if not result:
        fail("empty_file", "No bank transactions found")
    return {
        "source_name": filename,
        "source_sha256": hashlib.sha256(content).hexdigest(),
        "rows": result,
    }
