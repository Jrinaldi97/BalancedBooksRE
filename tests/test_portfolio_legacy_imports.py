"""Value-only legacy exports retain original byte and row identities."""

from datetime import datetime
import hashlib
import io
import struct
import pytest
import xlwt
from openpyxl import Workbook
from app.portfolio import models as m
from app.portfolio.imports import parse_bank
from app.portfolio.legacy_xls import _check_records
from app.portfolio.service import Problem
from tests.test_balancedbooks_portfolio import portfolio as _portfolio, run

portfolio = _portfolio

MAPPING = {"date": "Date", "amount": "Amount", "description": "Description"}


def workbook(kind="xls", sheets=("Bank",), formula=False, datemode=0):
    rows = [
        ["Date", "Amount", "Description", "Extra"],
        [datetime(2026, 1, 2), -1234.56, "Repeated fee", ""],
        [],
        [datetime(2026, 1, 2), -1234.56, "Repeated fee", ""],
        [datetime(2026, 1, 3), 45.67, "Deposit", ""],
    ]
    output = io.BytesIO()
    if kind == "xls":
        book = xlwt.Workbook()
        book.dates_1904 = datemode
        date_style = xlwt.easyxf(num_format_str="YYYY-MM-DD")
        for name in sheets:
            sheet = book.add_sheet(name)
            for number, row in enumerate(rows):
                for column, value in enumerate(row):
                    sheet.write(
                        number,
                        column,
                        value,
                        (
                            date_style
                            if isinstance(value, datetime)
                            else xlwt.Style.default_style
                        ),
                    )
            if formula:
                sheet.write(5, 3, xlwt.Formula("1+1"))
        book.save(output)
    else:
        book = Workbook()
        book.remove(book.active)
        for name in sheets:
            sheet = book.create_sheet(name)
            for row in rows:
                sheet.append(row)
            if formula:
                sheet.cell(6, 4, "=1+1")
        book.save(output)
    return output.getvalue()


def test_xls_xlsx_csv_exact_financial_and_provenance_parity():
    xls = workbook()
    xlsx = workbook("xlsx")
    csv = b"Date,Amount,Description,Extra\n2026-01-02,-1234.56,Repeated fee,\n\n2026-01-02,-1234.56,Repeated fee,\n2026-01-03,45.67,Deposit,\n"
    parsed = [
        parse_bank(data, "bank." + kind, MAPPING)
        for kind, data in (("xls", xls), ("xlsx", xlsx), ("csv", csv))
    ]
    expected = [
        {
            "date": "2026-01-02",
            "amount_cents": -123456,
            "description": "Repeated fee",
            "reference": "",
            "external_id": None,
        },
        {
            "date": "2026-01-02",
            "amount_cents": -123456,
            "description": "Repeated fee",
            "reference": "",
            "external_id": None,
        },
        {
            "date": "2026-01-03",
            "amount_cents": 4567,
            "description": "Deposit",
            "reference": "",
            "external_id": None,
        },
    ]
    for output, content in zip(parsed, (xls, xlsx, csv), strict=True):
        assert output["source_sha256"] == hashlib.sha256(content).hexdigest()
        assert [row["physical_row"] for row in output["rows"]] == [2, 4, 5]
        assert [
            {
                key: value
                for key, value in row.items()
                if key not in ("physical_row", "source_sheet")
            }
            for row in output["rows"]
        ] == expected
    assert [row["source_sheet"] for row in parsed[0]["rows"]] == ["Bank"] * 3
    assert (
        parse_bank(workbook(datemode=1), "bank.xls", MAPPING)["rows"]
        == parsed[0]["rows"]
    )


@pytest.mark.parametrize("kind", ["xls", "xlsx"])
def test_formulas_even_outside_mapping_are_rejected(kind):
    with pytest.raises(Problem) as caught:
        parse_bank(workbook(kind, formula=True), "bank." + kind, MAPPING)
    assert caught.value.code == "formula_cell"


@pytest.mark.parametrize("kind", ["xls", "xlsx"])
def test_sheets_are_explicit_and_same_file_rows_are_distinct(portfolio, kind):
    db, _, bank, _, _ = portfolio
    data = workbook(kind, sheets=("Checking", "Savings"))
    with pytest.raises(Problem) as caught:
        parse_bank(data, "bank." + kind, MAPPING)
    assert caught.value.code == "sheet_required"
    first = parse_bank(data, "bank." + kind, MAPPING, sheet="Checking")
    second = parse_bank(data, "bank." + kind, MAPPING, sheet="Savings")
    a = run(db, "import_bank", account_id=bank, **first)
    b = run(db, "import_bank", account_id=bank, **second)
    assert len(a["observation_ids"]) == len(b["observation_ids"]) == 3
    assert not set(a["observation_ids"]) & set(b["observation_ids"])
    assert run(db, "import_bank", account_id=bank, **first)["observation_ids"] == []
    rows = [db.get(m.BankObservation, ident) for ident in a["observation_ids"]]
    assert [row.source["occurrence"] for row in rows] == [1, 2, 1]
    assert [row.source["normalized_record_index"] for row in rows] == [1, 2, 3]
    assert [row.source["physical_row"] for row in rows] == [2, 4, 5]
    assert all(row.source["sha256"] == hashlib.sha256(data).hexdigest() for row in rows)


def test_legacy_same_sheet_import_keeps_existing_identity(portfolio):
    db, _, bank, _, _ = portfolio
    data = parse_bank(workbook("xlsx"), "bank.xlsx", MAPPING)
    result = run(db, "import_bank", account_id=bank, **data)
    for index, ident in enumerate(result["observation_ids"], 2):
        db.get(m.BankObservation, ident).source_key = f"{data['source_sha256']}:{index}"
    db.flush()
    again = run(db, "import_bank", account_id=bank, **data)
    assert again["observation_ids"] == []
    assert again["duplicate_ids"] == result["observation_ids"]


@pytest.mark.parametrize(
    "content,filename,code",
    [
        (b"Date,Amount\n2026-01-01,1", "bank.xls", "file_type"),
        (b"PK\x03\x04garbage", "bank.xlsx", "invalid_workbook"),
        (
            b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 512,
            "bank.xls",
            "invalid_workbook",
        ),
        (workbook(), "bank.csv", "file_type"),
        (workbook("xlsx"), "bank.xls", "file_type"),
    ],
)
def test_format_mismatch_and_corruption_fail_closed(content, filename, code):
    with pytest.raises(Problem) as caught:
        parse_bank(content, filename, MAPPING)
    assert caught.value.code == code


def test_encryption_formula_and_truncation_record_preflight():
    bof = struct.pack("<HHHH", 0x0809, 4, 0x0600, 5)
    eof = struct.pack("<HH", 0x000A, 0)
    _check_records(bof + eof + b"\0" * 100)
    for record, code in (
        (struct.pack("<HH", 0x002F, 0), "encrypted_workbook"),
        (struct.pack("<HH", 0x0006, 0), "formula_cell"),
        (struct.pack("<HH", 0x00FC, 500), "invalid_workbook"),
    ):
        with pytest.raises(Problem) as caught:
            _check_records(bof + record + eof)
        assert caught.value.code == code


def test_ambiguous_headers_missing_mapping_and_subcent_values_fail():
    for content, code in (
        (b"Date,Amount,Amount\n2026-01-01,1,2", "invalid_header"),
        (b"Date,Other\n2026-01-01,1", "missing_column"),
        (b"Date,Amount,Description\n2026-01-01,1.001,x", "invalid_amount"),
    ):
        with pytest.raises(Problem) as caught:
            parse_bank(content, "bank.csv", MAPPING)
        assert caught.value.code == code


@pytest.mark.parametrize(
    "value,style,code",
    [
        (True, None, "invalid_cell"),
        (1.001, None, "invalid_amount"),
        (60, "YYYY-MM-DD", "invalid_date"),
    ],
)
def test_xls_invalid_typed_cells_fail(value, style, code):
    book = xlwt.Workbook()
    sheet = book.add_sheet("Bank")
    for column, header in enumerate(["Date", "Amount", "Description"]):
        sheet.write(0, column, header)
    sheet.write(
        1,
        0,
        value if style else "2026-01-01",
        xlwt.easyxf(num_format_str=style) if style else xlwt.Style.default_style,
    )
    sheet.write(1, 1, 1 if style else value)
    output = io.BytesIO()
    book.save(output)
    with pytest.raises(Problem) as caught:
        parse_bank(output.getvalue(), "invalid.xls", MAPPING)
    assert caught.value.code == code


@pytest.mark.parametrize("part", ["xl/workbook.xml", "xl/worksheets/sheet1.xml"])
def test_xlsx_corrupt_xml_is_a_structured_rejection(part):
    from zipfile import ZipFile

    output = io.BytesIO()
    with ZipFile(io.BytesIO(workbook("xlsx"))) as original, ZipFile(
        output, "w"
    ) as broken:
        for name in original.namelist():
            broken.writestr(name, b"<broken>" if name == part else original.read(name))
    with pytest.raises(Problem) as caught:
        parse_bank(output.getvalue(), "corrupt.xlsx", MAPPING)
    assert caught.value.code == "invalid_workbook"
